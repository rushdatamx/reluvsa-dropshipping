"""Cruce Excel -> Excel de publicaciones con el catálogo vehicular de Autozur.

Este módulo NO conoce ni llama la API de Mercado Libre.  Lee los dos archivos que
entrega Autozur, propone compatibilidades a partir del título y deja cualquier
caso incompleto o ambiguo para revisión humana.
"""
from __future__ import annotations

import json
import re
import tempfile
import time
import unicodedata
import uuid
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import openpyxl


ENCABEZADOS_CATALOGO = (
    "FABRICANTE", "MODELO", "AÑO", "SUBMODELO", "LITROS", "CILINDROS",
    "CARROCERÍA", "TIPO DE TRANSMISIÓN", "TIPO DE TRACCIÓN",
    "TIPO DE COMBUSTIBLE", "TIPO DE MOTOR", "TIPO DE ASPIRACIÓN",
)

ENCABEZADOS_PUBLICACIONES = (
    "UserProductID", "TITULO", "SKU", *ENCABEZADOS_CATALOGO,
    "ASIGNACIÓN DE POSICIÓN(Conductor, Acompañante, Izquierda, Derecha, "
    "Delantera, Trasera, Interno, Externo, Superior, Inferior, Intermedio, Centro)",
    "NOTAS",
)

SESSION_DIR = Path(tempfile.gettempdir()) / "reluvsa_autozur"
SESSION_DIR.mkdir(parents=True, exist_ok=True)
SESSION_TTL_SEGUNDOS = 24 * 60 * 60

_ANIO = re.compile(r"\b(?:19|20)\d{2}\b")
_MOTOR = re.compile(r"\b([LVIHWR]\s*\d{1,2})\b", re.IGNORECASE)
_LITROS_CON_L = re.compile(r"\b(\d{1,2}(?:[.,]\d+)?)\s*[lL]\b")
_DECIMAL = re.compile(r"\b(\d{1,2}[.,]\d+)\b")
_MARCADOR = re.compile(r"\bP\s*/\s*", re.IGNORECASE)


def normalizar(valor) -> str:
    texto = "" if valor is None else str(valor).strip()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )
    texto = texto.upper().replace("–", "-").replace("—", "-")
    texto = re.sub(r"[^A-Z0-9./-]+", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def _texto_excel(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _encabezados(ws) -> Tuple[str, ...]:
    return tuple(_texto_excel(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1)))


def _validar_encabezados(actuales: Sequence[str], esperados: Sequence[str], nombre: str) -> None:
    if tuple(actuales[:len(esperados)]) == tuple(esperados):
        return
    faltantes = [h for h in esperados if h not in actuales]
    if faltantes:
        raise ValueError(f"El {nombre} no tiene las columnas requeridas: {', '.join(faltantes)}.")
    raise ValueError(f"El {nombre} contiene las columnas esperadas, pero no están en el orden de Autozur.")


def _anio_entero(valor) -> Optional[int]:
    texto = _texto_excel(valor)
    if not re.fullmatch(r"(?:19|20)\d{2}", texto):
        return None
    return int(texto)


def _numero_normalizado(valor) -> str:
    texto = _texto_excel(valor).replace(",", ".")
    try:
        numero = float(texto)
    except (TypeError, ValueError):
        return normalizar(texto)
    return f"{numero:g}"


@dataclass
class CatalogoIndexado:
    por_fabricante_modelo: Dict[Tuple[str, str], List[List[str]]]
    modelos_por_fabricante: Dict[str, List[str]]
    fabricantes_por_modelo: Dict[str, List[str]]
    fabricantes: List[str]
    filas: int


def leer_catalogo(ruta) -> CatalogoIndexado:
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb.active
    _validar_encabezados(_encabezados(ws), ENCABEZADOS_CATALOGO, "catálogo de vehículos")

    indice: Dict[Tuple[str, str], List[List[str]]] = defaultdict(list)
    modelos: Dict[str, set] = defaultdict(set)
    filas = 0
    for valores in ws.iter_rows(min_row=2, max_col=len(ENCABEZADOS_CATALOGO), values_only=True):
        fila = [_texto_excel(v) for v in valores]
        fabricante, modelo = normalizar(fila[0]), normalizar(fila[1])
        anio = _anio_entero(fila[2])
        if not fabricante or not modelo or anio is None:
            continue
        fila[2] = str(anio)
        indice[(fabricante, modelo)].append(fila)
        modelos[fabricante].add(modelo)
        filas += 1
    wb.close()
    return CatalogoIndexado(
        por_fabricante_modelo=dict(indice),
        modelos_por_fabricante={k: sorted(v, key=len, reverse=True) for k, v in modelos.items()},
        fabricantes_por_modelo={
            modelo: sorted(fabricantes)
            for modelo, fabricantes in _invertir_modelos(modelos).items()
        },
        fabricantes=sorted(modelos, key=len, reverse=True),
        filas=filas,
    )


def _invertir_modelos(modelos: Dict[str, set]) -> Dict[str, set]:
    resultado = defaultdict(set)
    for fabricante, nombres in modelos.items():
        for modelo in nombres:
            resultado[modelo].add(fabricante)
    return resultado


def leer_publicaciones(ruta) -> List[List[str]]:
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb.active
    _validar_encabezados(_encabezados(ws), ENCABEZADOS_PUBLICACIONES, "archivo de publicaciones")
    filas = []
    vistos = set()
    for numero, valores in enumerate(
        ws.iter_rows(min_row=2, max_col=len(ENCABEZADOS_PUBLICACIONES), values_only=True), start=2
    ):
        fila = [_texto_excel(v) for v in valores]
        if not any(fila):
            continue
        if not all(fila[i] for i in range(3)):
            raise ValueError(f"La fila {numero} necesita UserProductID, TITULO y SKU.")
        if fila[0] in vistos:
            raise ValueError(f"El UserProductID {fila[0]} está repetido en el archivo de entrada.")
        vistos.add(fila[0])
        filas.append(fila)
    wb.close()
    if not filas:
        raise ValueError("El archivo de publicaciones no contiene registros.")
    return filas


def _empieza_con(texto: str, opcion: str) -> bool:
    return texto == opcion or texto.startswith(opcion + " ")


def interpretar_titulo(titulo: str, catalogo: CatalogoIndexado) -> dict:
    limpio = normalizar(titulo)
    marcador = _MARCADOR.search(limpio)
    vehiculo = limpio[marcador.end():].strip() if marcador else limpio

    anios = [int(a) for a in _ANIO.findall(vehiculo)]
    anio_desde, anio_hasta = (min(anios), max(anios)) if anios else (None, None)

    motor_m = _MOTOR.search(vehiculo)
    cilindros = ""
    if motor_m:
        cilindros = re.sub(r"\s+", "", motor_m.group(1)).upper()

    litros_m = _LITROS_CON_L.search(vehiculo)
    if not litros_m:
        # Los títulos históricos suelen omitir la L: "Camaro 4.1 1975/1979".
        decimales = list(_DECIMAL.finditer(vehiculo))
        litros_m = decimales[-1] if decimales else None
    litros = _numero_normalizado(litros_m.group(1)) if litros_m else ""

    limite = len(vehiculo)
    cortes = [m.start() for m in (_ANIO.search(vehiculo), motor_m, litros_m) if m]
    if cortes:
        limite = min(cortes)
    cabeza = vehiculo[:limite].strip(" -/")

    alias_fabricante = {
        "VW": "VOLKSWAGEN",
        "MERCEDES BENZ": "MERCEDES-BENZ",
    }
    fabricante = ""
    fabricante_inferido = False
    modelo_ambiguo = False
    resto = cabeza
    opciones_fabricante = sorted(set(catalogo.fabricantes) | set(alias_fabricante), key=len, reverse=True)
    for opcion in opciones_fabricante:
        if _empieza_con(cabeza, opcion):
            fabricante = alias_fabricante.get(opcion, opcion)
            resto = cabeza[len(opcion):].strip(" -/")
            break

    modelo = ""
    if fabricante:
        for opcion in catalogo.modelos_por_fabricante.get(fabricante, []):
            if _empieza_con(resto, opcion):
                modelo = opcion
                break

    # Si el título omite fabricante (p. ej. "P/ Spark 1.2 2017"), sólo se puede
    # inferir cuando el modelo pertenece a una única marca. Si el mismo modelo
    # vive en varias marcas se conserva ambiguo.
    if not fabricante:
        modelos_globales = sorted(catalogo.fabricantes_por_modelo, key=len, reverse=True)
        for opcion in modelos_globales:
            if _empieza_con(cabeza, opcion):
                marcas = catalogo.fabricantes_por_modelo[opcion]
                if len(marcas) == 1:
                    fabricante, modelo = marcas[0], opcion
                    resto = cabeza
                    fabricante_inferido = True
                else:
                    modelo = opcion
                    modelo_ambiguo = True
                break

    detalle_modelo = resto[len(modelo):].strip(" -/") if modelo and _empieza_con(resto, modelo) else ""

    return {
        "fabricante": fabricante,
        "modelo": modelo,
        "litros": litros,
        "cilindros": cilindros,
        "anio_desde": anio_desde,
        "anio_hasta": anio_hasta,
        "texto_vehiculo": vehiculo,
        "modelo_detectado": resto,
        "detalle_modelo": detalle_modelo,
        "fabricante_inferido": fabricante_inferido,
        "modelo_ambiguo": modelo_ambiguo,
    }


def _filtrar_candidatos(parsed: dict, catalogo: CatalogoIndexado) -> List[List[str]]:
    fabricante, modelo = parsed["fabricante"], parsed["modelo"]
    if not fabricante or not modelo:
        return []
    candidatos = catalogo.por_fabricante_modelo.get((fabricante, modelo), [])
    if parsed["anio_desde"] is not None:
        candidatos = [
            f for f in candidatos
            if parsed["anio_desde"] <= int(f[2]) <= parsed["anio_hasta"]
        ]
    if parsed["litros"]:
        candidatos = [f for f in candidatos if _numero_normalizado(f[4]) == parsed["litros"]]
    if parsed["cilindros"]:
        numero = re.sub(r"\D", "", parsed["cilindros"])
        if numero:
            candidatos = [f for f in candidatos if _numero_normalizado(f[5]) == numero]

    unicos = []
    vistos = set()
    for fila in candidatos:
        llave = tuple(normalizar(v) for v in fila)
        if llave not in vistos:
            vistos.add(llave)
            unicos.append(fila)
    return unicos


def analizar(publicaciones: List[List[str]], catalogo: CatalogoIndexado) -> dict:
    resultados = []
    conteos = defaultdict(int)
    compatibilidades = 0

    for posicion, fila in enumerate(publicaciones):
        parsed = interpretar_titulo(fila[1], catalogo)
        candidatos = _filtrar_candidatos(parsed, catalogo)
        bloqueos = []
        observaciones = []
        if parsed["fabricante_inferido"]:
            observaciones.append(
                f"Fabricante inferido como {parsed['fabricante']} porque {parsed['modelo']} "
                "sólo corresponde a esta marca en el catálogo."
            )
        elif not parsed["fabricante"]:
            if parsed["modelo_ambiguo"]:
                bloqueos.append("El modelo aparece asociado a varias marcas en el catálogo.")
            else:
                bloqueos.append("No se identificó un fabricante al inicio de la sección P/.")
        if parsed["fabricante"] and not parsed["modelo"]:
            bloqueos.append("El modelo del título no coincide exactamente con el catálogo.")
        if parsed["anio_desde"] is None:
            bloqueos.append("El título no contiene un año utilizable.")
        if not parsed["litros"]:
            bloqueos.append("El título no contiene la cilindrada en litros.")
        if parsed["detalle_modelo"]:
            bloqueos.append(
                f"El título agrega «{parsed['detalle_modelo']}» después del modelo; revisa submodelo o carrocería."
            )
        if not candidatos and parsed["fabricante"] and parsed["modelo"]:
            propuestas = _filtrar_candidatos({**parsed, "litros": "", "cilindros": ""}, catalogo)
            if propuestas:
                candidatos = propuestas
                bloqueos.append(
                    "El motor exacto no aparece en el catálogo; se muestran opciones del mismo modelo y años."
                )
            elif not bloqueos:
                bloqueos.append("No se encontraron vehículos con la combinación exacta del título.")

        # La lista automática exige fabricante/modelo/años/cilindrada exactos.
        # Cilindros/configuración sólo restringen cuando aparecen explícitamente.
        # Las variantes de submodelo son candidatas válidas: se exportan todas.
        estado = "lista" if candidatos and not bloqueos else (
            "revision" if candidatos or parsed["modelo_ambiguo"] else "sin_coincidencia"
        )
        conteos[estado] += 1
        if estado == "lista":
            compatibilidades += len(candidatos)

        resultados.append({
            "id": posicion,
            "user_product_id": fila[0],
            "titulo": fila[1],
            "sku": fila[2],
            "estado": estado,
            "motivo": " ".join(observaciones + bloqueos),
            "extraido": parsed,
            "candidatos": candidatos,
        })

    return {
        "resultados": resultados,
        "resumen": {
            "publicaciones": len(publicaciones),
            "listas": conteos["lista"],
            "revision": conteos["revision"],
            "sin_coincidencia": conteos["sin_coincidencia"],
            "compatibilidades_automaticas": compatibilidades,
            "filas_catalogo": catalogo.filas,
        },
    }


def guardar_sesion(publicaciones: List[List[str]], analisis: dict) -> str:
    ahora = time.time()
    for antigua in SESSION_DIR.glob("*.json"):
        try:
            if ahora - antigua.stat().st_mtime > SESSION_TTL_SEGUNDOS:
                antigua.unlink(missing_ok=True)
        except OSError:
            pass
    session_id = uuid.uuid4().hex
    ruta = SESSION_DIR / f"{session_id}.json"
    payload = {"publicaciones": publicaciones, **analisis}
    ruta.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return session_id


def cargar_sesion(session_id: str) -> dict:
    if not re.fullmatch(r"[0-9a-f]{32}", session_id or ""):
        raise ValueError("La sesión de análisis no es válida.")
    ruta = SESSION_DIR / f"{session_id}.json"
    if not ruta.exists():
        raise ValueError("La sesión de análisis expiró. Vuelve a cargar los dos archivos.")
    return json.loads(ruta.read_text(encoding="utf-8"))


def resumen_para_cliente(session_id: str, analisis: dict) -> dict:
    items = []
    for r in analisis["resultados"]:
        items.append({
            k: r[k] for k in ("id", "user_product_id", "titulo", "sku", "estado", "motivo", "extraido")
        } | {
            "compatibilidades": len(r["candidatos"]),
            "muestra": r["candidatos"][:5],
        })
    return {"session_id": session_id, **analisis["resumen"], "resultados": items}


def escribir_resultado(session: dict, aprobadas_revision: Iterable[int], excluidas: Iterable[int], destino) -> int:
    aprobadas = {int(i) for i in aprobadas_revision}
    excluidas_ids = {int(i) for i in excluidas}
    resultados = {r["id"]: r for r in session["resultados"]}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compatibilidades"
    ws.append(list(ENCABEZADOS_PUBLICACIONES))

    total = 0
    for posicion, original in enumerate(session["publicaciones"]):
        resultado = resultados[posicion]
        incluir = posicion not in excluidas_ids and (resultado["estado"] == "lista" or (
            resultado["estado"] == "revision" and posicion in aprobadas
        ))
        if not incluir:
            continue
        for compatibilidad in resultado["candidatos"]:
            # Autozur pide una fila por compatibilidad. Se repiten los tres datos
            # de publicación y se conservan posición/notas del archivo original.
            ws.append([
                original[0], original[1], original[2], *compatibilidad,
                original[15], original[16],
            ])
            total += 1

    if not total:
        raise ValueError("No hay compatibilidades aprobadas para generar el archivo.")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for columna, ancho in {"A": 18, "B": 62, "C": 18, "D": 20, "E": 24, "F": 10,
                           "G": 24, "H": 10, "I": 11, "J": 18, "P": 28, "Q": 28}.items():
        ws.column_dimensions[columna].width = ancho
    wb.save(destino)
    return total
