"""Parser de la plantilla ERP de conciliación de facturas.

La plantilla no se identifica por nombre: exige Referencia, Proveedor -> Nombre,
Fecha Referencia y Fecha Captura. Conserva cada fila aunque tenga una anomalía
para que el administrador pueda revisarla en la conciliación.
"""
from datetime import date, datetime
import json
import re
import unicodedata
from pathlib import Path

import openpyxl
from fastapi import HTTPException


def _clave(value) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return " ".join(text.lower().split())


def _texto(value) -> str:
    if value is None:
        return ""
    # Excel suele leer folios como float; no debemos inventar el .0.
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _fecha(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _texto(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def _referencia(value) -> str:
    return re.sub(r"[\s-]+", "", _texto(value)).upper()


def _proveedor(nombre: str):
    aliases = {
        "kims auto corporation": "KIM",
        "quality hoses sa de cv": "CAUPLAS",
    }
    return aliases.get(_clave(nombre))


def parse_erp(path: Path):
    """Devuelve las entradas auditables y el rango de Fecha Referencia.

    Sólo las fechas de referencia son obligatorias para aceptar la plantilla: una
    fila individual errónea no se pierde, queda marcada ``Para revisar``.
    """
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="No se pudo abrir el archivo Excel") from exc

    try:
        ws = wb[wb.sheetnames[0]]
        header_row = None
        headers = {}
        required = {"referencia", "fecha referencia", "fecha captura"}
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            candidate = {_clave(v): index for index, v in enumerate(row) if _clave(v)}
            # "Proveedor -> Nombre" puede llegar con flecha, salto de línea o sólo Nombre.
            proveedor_idx = next((i for k, i in candidate.items() if "proveedor" in k and "nombre" in k), None)
            if proveedor_idx is None:
                proveedor_idx = candidate.get("nombre")
            if required.issubset(candidate) and proveedor_idx is not None:
                header_row = row_idx
                headers = candidate
                headers["__proveedor__"] = proveedor_idx
                break
        if header_row is None:
            raise HTTPException(
                status_code=400,
                detail="El Excel ERP debe incluir Referencia, Proveedor -> Nombre, Fecha Referencia y Fecha Captura.",
            )

        entries = []
        fechas = []
        for excel_row, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not any(v is not None and _texto(v) for v in row):
                continue
            raw = {str(i): _texto(value) for i, value in enumerate(row)}
            proveedor_original = _texto(row[headers["__proveedor__"]])
            referencia_original = _texto(row[headers["referencia"]])
            fecha_referencia = _fecha(row[headers["fecha referencia"]])
            fecha_captura = _fecha(row[headers["fecha captura"]])
            codigo = _proveedor(proveedor_original)
            ref = _referencia(referencia_original)
            motivos = []
            if not codigo:
                motivos.append("Proveedor no reconocido en ERP")
            if not ref:
                motivos.append("Referencia vacía o inválida")
            elif codigo == "KIM" and not ref.startswith("K"):
                motivos.append("La referencia KIMS debe iniciar con K")
            elif codigo == "CAUPLAS" and not ref.isdigit():
                motivos.append("La referencia CAUPLAS debe ser el folio numérico")
            if not fecha_referencia:
                motivos.append("Fecha Referencia requerida o inválida")
            if not fecha_captura:
                motivos.append("Fecha Captura requerida o inválida")
            if fecha_referencia:
                fechas.append(fecha_referencia)
            entries.append({
                "numero_fila": excel_row,
                "fila_original_json": json.dumps(raw, ensure_ascii=False),
                "proveedor_original": proveedor_original,
                "proveedor_codigo": codigo,
                "referencia_original": referencia_original,
                "referencia_normalizada": ref or None,
                "fecha_referencia": fecha_referencia,
                "fecha_captura": fecha_captura,
                "motivo_revision": "; ".join(motivos) or None,
            })
        if not entries:
            raise HTTPException(status_code=400, detail="El Excel ERP no contiene filas de datos.")
        if not fechas:
            raise HTTPException(status_code=400, detail="El Excel ERP no contiene ninguna Fecha Referencia válida.")
        return {"entradas": entries, "fecha_desde": min(fechas), "fecha_hasta": max(fechas)}
    finally:
        wb.close()
