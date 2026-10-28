"""Regresión del master DIFIRSA de Publicaciones Masivas."""
import os
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from services.generador_plantilla import COLUMNAS, ConfiguracionProveedor, escribir_xlsx, generar_filas_con_reporte
from services.parser_catalogo import leer_catalogo_detallado
from services.perfiles_catalogo import perfil_de, proveedores_soportados

assert "DIFIRSA" in proveedores_soportados()
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Master"
ws.append(["Producto", "Armadora", "Modelo", "Motor", "Inicio", "Fin", "Notas", "Engine Fil",
           "Eq1", "Eq2", "Eq3", "Eq4", "Eq5", "Eq6", "Eq7", "OEM", "Precio", "Stock"])
ws.append(["Bomba de agua", "Ford", "F-150", "V8 5.0L", 2015, 2018, "", "EF-123",
           "EQ-A", "EQ-B", "", "", "", "", "", "OEM-1", 100, None])
ws.append(["Bomba de agua", "Ford", "", "", 2015, None, "", "EF-123",
           "", "", "", "", "", "", "", "OEM-2", "#N/A", None])
otra = wb.create_sheet("No procesar")
otra.append(["Producto", "Armadora", "Modelo", "Motor", "Inicio", "Fin", "Notas", "Engine Fil"])
otra.append(["Pieza ajena", "Chevrolet", "Cruze", "L4", 2010, 2012, "", "NO-LEER"])
tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); tmp.close(); wb.save(tmp.name)
lectura = leer_catalogo_detallado(tmp.name, perfil_de("DIFIRSA")); os.unlink(tmp.name)
assert lectura.formato == "master_difrisa" and lectura.filas_master == 2
assert len(lectura.piezas) == 1 and lectura.piezas[0]["clave"] == "EF-123"
pieza = lectura.piezas[0]
assert pieza["costo"] is None and pieza["stock"] is None
assert pieza["equivalencias"] == ["EQ-A", "EQ-B"] and pieza["oems"] == ["OEM-1", "OEM-2"]
assert lectura.compatibilidades_validas == 1 and any("incompleta" in e["motivo"] for e in lectura.errores)
assert lectura.sku_sin_precio == 1
filas, _ = generar_filas_con_reporte(lectura.piezas, ConfiguracionProveedor(codigo_bodega="DIFIRSA"))
assert len(filas) == 3 and all(f.marca == "Engine Fil" and f.precio is None for f in filas)
assert "EQ-A" in filas[0].descripcion and "OEM-2" in filas[0].descripcion
assert all(f.stock is None and f.imagenes == [] for f in filas)
assert len(COLUMNAS) == 38 and COLUMNAS[37] == "DIFIRSA"
out = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); out.close()
escribir_xlsx(filas[:1], ConfiguracionProveedor(codigo_bodega="DIFIRSA"), out.name)
salida = openpyxl.load_workbook(out.name, data_only=True); os.unlink(out.name)
hoja = salida.active
assert [hoja.cell(1, i).value for i in range(1, 38)] == COLUMNAS[:37]
assert hoja.cell(1, 38).value == "DIFIRSA"
assert hoja.cell(2, COLUMNAS.index("Precio") + 1).value is None
assert hoja.cell(2, COLUMNAS.index("Envio Gratis(si,no)") + 1).value is None
assert hoja.cell(2, COLUMNAS.index("Cantidad") + 1).value is None
assert hoja.cell(2, 38).value is None
assert all(hoja.cell(2, COLUMNAS.index(f"Imagen{i}") + 1).value is None for i in range(1, 11))

# Formato vigente de Gaby: litros en E y año simple/rango en F.
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "FILTROS DE AIRE"
ws.append(["Producto", "Armadora", "Automóvil", "Cilindros", "Litros", "Año", "Notas", "Engine Fil",
           "Equiv 1", "Equiv 2", "Equiv 3", "Equiv 4", "Equiv 5", "Equiv 6", "Equiv 7", "OEM", "Precio", "Stock"])
ws.append(["Filtro de aire", "Ford", "Focus", "L4", 2.5, "2007-2010", "", "EF-NEW", "ALT-1", "", "", "", "", "", "", "OEM-X", 245.5, 12])
ws.append(["Filtro de aire", "Ford", "Fiesta", "L4", 1.6, 2014, "", "EF-NEW", "", "", "", "", "", "", "", "OEM-Y", 245.5, 12])
ws.append(["Filtro de aire", "Ford", "Focus", "L4", 2.0, "20XX", "", "EF-NEW", "", "", "", "", "", "", "", "OEM-Z", 245.5, 12])
ws.append(["Filtro de aire", "Ford", "", "L4", 2.0, 2016, "", "EF-NEW", "", "", "", "", "", "", "", "OEM-W", 245.5, 12])
ws.append(["Filtro de aire", "Ford", "Focus", "L4", 2.0, 2016, "", "", "", "", "", "", "", "", "", "OEM-W", 245.5, 12])
ws.append(["", "Ford", "Focus", "L4", 2.0, 2016, "", "EF-MISSING", "", "", "", "", "", "", "", "", 10, 1])
tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); tmp.close(); wb.save(tmp.name)
lectura = leer_catalogo_detallado(tmp.name, perfil_de("DIFIRSA")); os.unlink(tmp.name)
assert lectura.formato == "master_difrisa" and lectura.filas_master == 6
assert len(lectura.piezas) == 1 and lectura.piezas[0]["clave"] == "EF-NEW"
pieza = lectura.piezas[0]
assert pieza["producto"] == "Filtro de aire" and pieza["costo"] == 245.5 and pieza["stock"] == 12
assert pieza["equivalencias"] == ["ALT-1"] and pieza["oems"] == ["OEM-X", "OEM-Y", "OEM-Z", "OEM-W"]
assert [(c["inicio"], c["fin"], c["motor"]) for c in pieza["compatibilidades"]] == [
    (2007, 2010, "L4 2.5L"), (2014, 2014, "L4 1.6L")]
assert lectura.compatibilidades_validas == 2 and lectura.compatibilidades_invalidas == 2
assert any("Falta Engine Fil o Producto" in e["motivo"] for e in lectura.errores)
print("DIFIRSA: OK")
