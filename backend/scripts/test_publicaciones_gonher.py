"""Regresión autocontenida del master GONHER."""
import os
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from services.generador_plantilla import (COLUMNAS, ConfiguracionProveedor,
    escribir_xlsx, generar_filas_con_reporte)
from services.parser_catalogo import leer_catalogo_detallado
from services.perfiles_catalogo import perfil_de

TOTAL = OK = 0
def check(nombre, condicion, detalle=None):
    global TOTAL, OK
    TOTAL += 1
    if condicion:
        OK += 1; print("  OK", nombre)
    else:
        print("FALLA", nombre, detalle or "")

HEADERS = ["Producto", "Línea", "Código Actual", "Filtro", "Auxiliar", "Marca", "Modelo",
           "Motor", "Año", "Tipo", "Rosca", "Altura cm", "Ø Ext. cm", "Ø Int. cm",
           "Ø Int. cm", "Largo. cm", "Ancho. cm", "OEM 1", "OEM 2", "Parte FRAM",
           "Parte Interfill", "Código de barras ", "Código SAT", "Precio", "Stock"]

def fila(sku="G-1048", marca="Ford", modelo="F-250 Super Duty", motor="V8 6.4L Turbo diésel",
         anio="2008-2010", precio=900.14, linea="GASOLINA"):
    return ["Filtro", linea, 1000201, sku, 1, marca, modelo, motor, anio, "Cartucho", "^^",
            12.11, 9.07, 2.81, "^^", "- NA -", "N/A", "8C3Z-9N184-A", "8C3Z-9N184-C",
            "CS10263A", "FGI-338D", 744926014498, 40161513, precio, None]

def guardar(wb):
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); tmp.close(); wb.save(tmp.name)
    return tmp.name

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Catálogo renombrado"; ws.append(HEADERS)
ws.append(fila()); ws.append(fila(modelo="F-350 Super Duty")); ws.append(fila(anio="2019-2017"))
ws.append(fila(sku="SIN-PRECIO", precio="ND", linea="AIRE"))
gc = wb.create_sheet("GC"); gc.append(HEADERS); gc.append(fila(sku="NO-LEER"))
ruta = guardar(wb); r = leer_catalogo_detallado(ruta, perfil_de("GONHER")); os.unlink(ruta)
check("detecta GONHER por encabezados aunque cambie la hoja", r.formato == "master_gonher")
check("ignora GC", r.sku_unicos_master == 2)
check("una aplicación inválida no elimina el SKU", r.compatibilidades_invalidas == 1 and len(r.piezas) == 2)
pieza = next(p for p in r.piezas if p["clave"] == "G-1048")
check("conserva ambos diámetros por ocurrencia", pieza["medidas"]["Diámetro interior 1"] == ["2.81"] and not pieza["medidas"]["Diámetro interior 2"])
check("omite placeholders", not pieza["roscas"] and not pieza["medidas"]["Largo"])
check("costo inválido queda vacío y advertido", r.sku_sin_precio == 1 and next(p for p in r.piezas if p["clave"] == "SIN-PRECIO")["costo"] is None)

cfg = ConfiguracionProveedor("GONHER", descripcion_base="BASE RELUVSA", marca="GONHER")
filas, reporte = generar_filas_con_reporte(r.piezas, cfg)
primera = next(f for f in filas if f.sku == "G-1048")
check("título reduce motor antes de quitar marca", primera.titulo == "Filtro de gasolina P/ F-250 Super Duty V8 6.4L 2008-2010", primera.titulo)
check("descripción consolida sin barras ni SAT", all(x in primera.descripcion for x in ("Código actual: 1000201", "Altura: 12.11 cm", "OEM: 8C3Z-9N184-A | 8C3Z-9N184-C", "FRAM: CS10263A", "Interfill: FGI-338D", "BASE RELUVSA")) and "744926" not in primera.descripcion and "40161513" not in primera.descripcion)
check("precio válido se calcula y el inválido queda vacío", primera.precio is not None and next(f for f in filas if f.sku == "SIN-PRECIO").precio is None)

titulos = openpyxl.Workbook(); ts = titulos.active; ts.append(HEADERS)
ts.append(fila(sku="FULL", modelo="Fiesta", motor="L4 1.6L", anio="2018"))
ts.append(fila(sku="RED", modelo="F-250 Super Duty", motor="V8 6.4L Turbo diésel", anio="2008-2010"))
ts.append(fila(sku="SIN-MARCA", marca="Marca Excesivamente Larga", modelo="Modelo Mediano", motor="V8 6.4L Turbo", anio="2010-2012"))
ts.append(fila(sku="FUERA", marca="Marca Muy Larga", modelo="Modelo Cuya Denominación Es Definitivamente Demasiado Larga", motor="V8 6.4L Turbo", anio="2010-2012"))
ts.append(fila(sku="FULL", modelo="Fiesta", motor="L4 1.6L", anio="2018"))
ruta = guardar(titulos); rt = leer_catalogo_detallado(ruta, perfil_de("GONHER")); os.unlink(ruta)
ft, reportet = generar_filas_con_reporte(rt.piezas, cfg)
check("título completo se conserva cuando cabe", next(f for f in ft if f.sku == "FULL").titulo == "Filtro de gasolina P/ Ford Fiesta L4 1.6L 2018")
check("tercer intento omite marca", "Marca Excesivamente" not in next(f for f in ft if f.sku == "SIN-MARCA").titulo)
check("título imposible se excluye sin recortar", any(e["clave"] == "FUERA" for e in reportet["exclusiones"]) and all(len(f.titulo) <= 60 for f in ft))
check("deduplica SKU + título normalizado", sum(f.sku == "FULL" for f in ft) == 1)

salida = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); salida.close()
escribir_xlsx(filas, cfg, salida.name); wo = openpyxl.load_workbook(salida.name, data_only=True); so = wo.active
idx = {c.value: c.column for c in so[1]}
check("plantilla tiene 37 columnas y GONHER al final", so.max_column == 37 and COLUMNAS[-1] == "GONHER")
check("stock GONHER queda vacío y otras bodegas en cero", so.cell(2, idx["Cantidad"]).value is None and so.cell(2, idx["GONHER"]).value is None and all(so.cell(2, idx[b]).value == 0 for b in ("AG", "CAUPLAS", "KG", "KIM", "MATRIZ", "VAZLO")))
os.unlink(salida.name)

incompleto = openpyxl.Workbook(); incompleto.active.append([h for h in HEADERS if h != "Precio"])
ruta = guardar(incompleto)
try:
    leer_catalogo_detallado(ruta, perfil_de("GONHER")); mensaje = ""
except ValueError as exc:
    mensaje = str(exc)
os.unlink(ruta)
check("huella incompleta enumera faltantes", "Precio" in mensaje and "faltan" in mensaje)

print(f"\nGONHER: {OK}/{TOTAL}")
if OK != TOTAL: raise SystemExit(1)
