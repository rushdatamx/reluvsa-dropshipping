"""Regresión del master KIMS en Publicaciones masivas."""
import os
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from services.generador_plantilla import (COLUMNAS, ConfiguracionProveedor,
    escribir_xlsx, generar_filas_con_reporte)
from services.parser_catalogo import KIMS_HEADERS, leer_catalogo_detallado
from services.perfiles_catalogo import perfil_de
from services.precio_publicacion import ParametrosPrecio, calcular_precio

ok = total = 0
def check(nombre, condicion, detalle=""):
    global ok, total
    total += 1
    if condicion:
        ok += 1; print(f"  ✅ {nombre}")
    else:
        print(f"  ❌ {nombre}: {detalle}")


def guardar(filas, headers=KIMS_HEADERS):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Renombrada"
    ws.append(list(headers))
    for fila in filas: ws.append([fila.get(h) for h in headers])
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); tmp.close(); wb.save(tmp.name)
    return tmp.name


base = {h: "" for h in KIMS_HEADERS}
base.update({"original": 12345.0, "alterna_1": "OEM-A", "alterna_2": "OEM-A",
    "armadora": "Ford", "modelo": "F-150", "version": "XL", "Año Inicio": 2018,
    "Año final": 2020, "motor": "V8 5.0L", "combustible": "Gasolina",
    "comentarios": "Incluye empaque", "posicion_1": "Delantera", "descripcion": "BOMBA DE AGUA",
    "cantidad": 7.0, "marca": "KIMS PRO", "sistema": "Enfriamiento", "PRECIO": 10,
    "FOTO 1": "https://www.kimsauto.com.mx/a.png", "FOTO 3": "https://www.kimsauto.com.mx/c.png"})
duplicada = dict(base)
sin_motor = dict(base, **{"armadora": "Nissan", "modelo": "NP300", "version": "", "motor": "",
                           "Año Inicio": 2021, "Año final": 2021, "sistema": "Motor"})
anios_malos = dict(base, **{"Año Inicio": 2025, "Año final": 2020})
ruta = guardar([base, duplicada, sin_motor, anios_malos])
lectura = leer_catalogo_detallado(ruta, perfil_de("KIM")); os.unlink(ruta)

check("detecta encabezados normalizados antes de otros formatos", lectura.formato == "master_kims")
check("SKU numérico pierde .0", lectura.piezas[0]["clave"] == "12345", lectura.piezas[0]["clave"])
check("agrupa filas por original", len(lectura.piezas) == 1)
check("deduplica compatibilidad repetida", lectura.duplicados_descartados == 1)
check("excluye sólo compatibilidad con años invertidos", lectura.compatibilidades_invalidas == 1 and len(lectura.piezas[0]["compatibilidades"]) == 2)
check("motor vacío sigue siendo utilizable", lectura.piezas[0]["compatibilidades"][1]["motor"] == "")
check("consolida OEM B:E sin duplicados", lectura.piezas[0]["oems"] == ["OEM-A"], lectura.piezas[0]["oems"])

cfg = ConfiguracionProveedor("KIM", descripcion_base="BASE RELUVSA", tipo_cambio_usd=18.5,
    params_precio=ParametrosPrecio(iva=.16, utilidad=.5, comision_ml=.13, envio_default=0))
filas, reporte = generar_filas_con_reporte(lectura.piezas, cfg)
check("rango 2018–2020 genera tres años y el rango completo", len(filas) == 5,
      [f.titulo for f in filas])
check("títulos no exceden 60", all(len(f.titulo) <= 60 for f in filas))
check("títulos KIM usan para", all(" para " in f.titulo for f in filas))
titulos_anuales = [f.titulo for f in filas if f.aplicacion in {"2018", "2019", "2020"}]
check("títulos anuales conservan producto modelo motor y año", len(titulos_anuales) == 3 and
      all("Bomba de Agua" in t and "F-150" in t and "V8 5.0L" in t and
          any(t.endswith(" " + a) for a in ("2018", "2019", "2020")) for t in titulos_anuales), titulos_anuales)
check("incluye además el título con rango completo",
      any(f.aplicacion == "2018/2020" and f.titulo.endswith("2018/2020") for f in filas))
check("rango de un año produce una sola variante",
      sum(f.aplicacion == "2021" for f in filas) == 1,
      [f.aplicacion for f in filas])
check("descripción consolida datos y base", all(x in filas[0].descripcion for x in ("OEM-A", "Gasolina", "Incluye empaque", "Delantera", "BASE RELUVSA", "Nissan NP300 2021")))
esperado = calcular_precio(10 * 18.5, "Bomba De Agua", cfg.params_precio)
check("convierte USD a MXN antes de fórmula compartida", filas[0].precio == esperado, (filas[0].precio, esperado))
check("marca propia viaja por fila", filas[0].marca == "KIMS PRO")

tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); tmp.close()
escribir_xlsx(filas, cfg, tmp.name)
wb = openpyxl.load_workbook(tmp.name, data_only=True); ws = wb.active
check("XLSX conserva las 37 columnas", ws.max_column == 37 and [c.value for c in ws[1]] == COLUMNAS)
indices = {c.value: c.column for c in ws[1]}
check("XLSX escribe cada una de las cinco variantes", ws.max_row == 6, ws.max_row)
check("XLSX escribe marca del catálogo en cada variante",
      all(ws.cell(r, indices["Marca"]).value == "KIMS PRO" for r in range(2, ws.max_row + 1)))
check("stock sólo va en KIM en cada variante", all(
      ws.cell(r, indices["Cantidad"]).value == 7 and ws.cell(r, indices["KIM"]).value == 7 and
      all(ws.cell(r, indices[b]).value == 0 for b in ("AG", "CAUPLAS", "KG", "MATRIZ", "VAZLO", "GONHER"))
      for r in range(2, ws.max_row + 1)))
check("fotos conservan posición sin desplazarse en cada variante", all(
      ws.cell(r, indices["Imagen1"]).value.endswith("a.png") and
      ws.cell(r, indices["Imagen2"]).value is None and
      ws.cell(r, indices["Imagen3"]).value.endswith("c.png") for r in range(2, ws.max_row + 1)))
os.unlink(tmp.name)

conflictos = []
for cambio in ({"original": "BAD-STOCK", "cantidad": 0},
               {"original": "BAD-MARCA", "marca": ""},
               {"original": "BAD-PRODUCTO", "descripcion": ""}):
    conflictos.append(dict(base, **cambio))
ruta = guardar(conflictos); malos = leer_catalogo_detallado(ruta, perfil_de("KIM")); os.unlink(ruta)
check("excluye SKU completos por stock/precio, marca y producto", len(malos.piezas) == 0 and malos.metricas_kims == {
    "sku_excluidos_precio_stock": 1, "sku_excluidos_marca": 1,
    "sku_excluidos_producto": 1, "sku_excluidos_estructura": 0}, malos.metricas_kims)

largo = dict(base, **{"original": "LARGO", "descripcion": "Producto Extraordinariamente Largo Imposible",
                      "modelo": "Modelo Con Nombre Demasiado Largo", "motor": "V8 Motor Intocable"})
ruta = guardar([largo]); lr = leer_catalogo_detallado(ruta, perfil_de("KIM")); os.unlink(ruta)
lf, rr = generar_filas_con_reporte(lr.piezas, cfg)
check("variante imposible no se corta y se reporta", not lf and rr["exclusiones"][0]["motivo"] == "Título excede 60 caracteres")

print(f"\n{ok}/{total} pruebas KIMS correctas")
if ok != total: raise SystemExit(1)
