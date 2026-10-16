"""Flujo postventa y API interna sin red ni base de producción."""
import os
import sys
import tempfile
from pathlib import Path

tmp = tempfile.TemporaryDirectory()
os.environ["DATABASE_PATH"] = str(Path(tmp.name) / "test.db")
sys.path.insert(0, str(Path(__file__).parent.parent))

from fastapi.testclient import TestClient  # noqa: E402
from jose import jwt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from database import get_db  # noqa: E402
from main import app  # noqa: E402
from models import UserInfo  # noqa: E402
from routers.auth import ALGORITHM, SECRET_KEY, require_admin  # noqa: E402
from services import ml_client, sync_devoluciones, sync_ml  # noqa: E402
from services.parser_albaran import parse_albaran  # noqa: E402

with get_db() as conn:
    conn.execute("INSERT INTO ml_config (clave, valor) VALUES ('seller_id', 'seller-1')")
    conn.execute("""INSERT INTO ventas_ml
        (num_venta, num_venta_origen, cuenta_ml, sku, deposito, pack_id)
        VALUES ('order-1','order-1','principal','SKU-1','KIM','pack-1')""")
    conn.execute("""INSERT INTO ventas_ml
        (num_venta, num_venta_origen, cuenta_ml, sku, deposito, pack_id)
        VALUES ('secundaria::order-1','order-1','secundaria','WRONG','AG','pack-1')""")
    for resource in ("/post-purchase/v1/claims/101", "/post-purchase/v1/claims/101",
                     "/post-purchase/v1/claims/102", "/post-purchase/v1/claims/103"):
        conn.execute("""INSERT INTO ml_notificaciones
            (topic, resource, recibido_en, cuenta_ml) VALUES ('post_purchase', ?, '2026-10-16', 'principal')""", (resource,))

claims = {
    "101": {"id": 101, "resource": "order", "resource_id": "order-1", "status": "opened",
            "type": "return", "reason_id": "PDD1", "claimed_quantity": 2,
            "last_updated": "2026-10-16T12:00:00+00:00",
            "players": [{"role": "respondent", "type": "seller", "available_actions": [
                {"action": "review", "mandatory": True, "due_date": "2026-10-17T12:00:00+00:00"}]}]},
    "102": {"id": 102, "resource": "order", "resource_id": "order-1", "status": "opened", "type": "mediations"},
    "103": {"id": 103, "resource": "order", "resource_id": "missing", "status": "opened", "type": "return"},
    "104": {"id": 104, "resource": "order", "resource_id": "order-1", "status": "closed", "type": "returns"},
}
search_params = []
optional_paths = []
fail = {"status": "403"}
old_get, old_optional = ml_client.get, ml_client.get_opcional


def fake_get(path, **kwargs):
    if path == "/post-purchase/v1/claims/search":
        search_params.append(kwargs.get("params") or {})
        return {"data": [], "paging": {"total": 0}}
    if path.endswith("/103") and fail["status"]:
        raise ml_client.MLError("GET falló " + fail["status"])
    if path.startswith("/orders/"):
        return {"order_items": [{"item": {"seller_sku": "SKU-1", "title": "Uno"}, "quantity": 1},
                                {"item": {"seller_sku": "SKU-2", "title": "Dos"}, "quantity": 2}]}
    return claims[path.rsplit("/", 1)[-1]]


def fake_optional(path, **kwargs):
    optional_paths.append(path)
    if path.endswith("/detail"):
        return {"due_date": "2026-10-17T12:00:00+00:00", "action_responsible": "seller" if "/101/" in path else "mediator"}
    if "/reasons/" in path:
        return {"detail": "Producto dañado"}
    if path.endswith("/returns"):
        return {"id": 55, "status": "in_progress"} if "/101/" in path else None
    if path.startswith("/orders/"):
        return fake_get(path)
    return None


ml_client.get, ml_client.get_opcional = fake_get, fake_optional
try:
    result = sync_devoluciones.procesar_pendientes()
    assert result == {"procesados": 2, "errores": 1}, result
    assert not any(path.endswith("/102/detail") for path in optional_paths)
    with get_db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM ml_devoluciones").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM ml_postventa_sin_devolucion").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM ml_notificaciones WHERE procesada=0").fetchone()[0] == 1
    fail["status"] = "429"
    sync_devoluciones._retry_at.clear()  # simula que transcurrió el backoff
    assert sync_devoluciones.procesar_pendientes()["errores"] == 1
    fail["status"] = None
    sync_devoluciones._retry_at.clear()  # simula que transcurrió el backoff
    assert sync_devoluciones.procesar_pendientes()["procesados"] == 1
    app.dependency_overrides[require_admin] = lambda: UserInfo(user_id=1, email="admin@test.mx", rol="admin", proveedor_id=None)
    client = TestClient(app)
    listing = client.get("/api/devoluciones").json()
    assert [i["claim_id"] for i in listing["items"]] == ["101", "103"]
    first = listing["items"][0]
    assert first["sku"] == "SKU-1" and first["unidades_reclamadas"] == 2
    assert len(first["productos"]) == 2 and first["fecha_limite"]
    assert listing["items"][1]["sku"] is None and listing["items"][1]["fecha_limite"] is None
    assert client.get("/api/devoluciones", params={"almacen": "KIM"}).json()["total"] == 1
    changed = client.patch("/api/devoluciones/principal/101", json={"mediacion": True, "revisado": True}).json()
    assert changed["mediacion"] and changed["revisado"] and changed["seguimiento_usuario"] == "admin@test.mx"
    sync_devoluciones.procesar_reclamo("principal", "101")
    assert client.get("/api/devoluciones/principal/101").json()["revisado"] == 1
    assert sync_devoluciones.procesar_reclamo("principal", "104") is True
    assert not any(path.endswith("/104/detail") for path in optional_paths)
    bad = client.post("/api/devoluciones/principal/101/evidencias", files={"file": ("bad.jpg", b"bad", "image/jpeg")})
    assert bad.status_code == 400
    good = client.post("/api/devoluciones/principal/101/evidencias", files={"file": ("photo.jpg", b"\xff\xd8\xffabc\xff\xd9", "image/jpeg")})
    assert good.status_code == 200, good.text
    fid = client.get("/api/devoluciones/principal/101").json()["evidencias"][0]["id"]
    assert client.get(f"/api/devoluciones/principal/101/evidencias/{fid}").status_code == 200
    huge = client.post("/api/devoluciones/principal/101/evidencias",
                       files={"file": ("huge.pdf", b"%PDF-" + b"x" * (10 * 1024 * 1024), "application/pdf")})
    assert huge.status_code == 413
    with get_db() as conn:
        for n in range(9):
            conn.execute("""INSERT INTO ml_devolucion_archivos
                (cuenta_ml,claim_id,nombre_original,nombre_servidor,tipo,tamano,usuario,creado_en)
                VALUES ('principal','101','x.pdf',?,'application/pdf',1,'admin@test.mx','2026-10-16')""", (f"test-{n}.pdf",))
    assert client.post("/api/devoluciones/principal/101/evidencias",
                       files={"file": ("last.pdf", b"%PDF-abc%%EOF", "application/pdf")}).status_code == 400
    app.dependency_overrides.clear()
    assert TestClient(app).get("/api/devoluciones").status_code in (401, 403)
    token = jwt.encode({"sub": "2", "email": "proveedor@test.mx", "rol": "proveedor",
                        "proveedor_id": 1,
                        "exp": __import__('datetime').datetime.now(__import__('datetime').timezone.utc) +
                        __import__('datetime').timedelta(hours=1)}, SECRET_KEY, algorithm=ALGORITHM)
    assert TestClient(app).get("/api/devoluciones", headers={"Authorization": f"Bearer {token}"}).status_code == 403
    with get_db() as conn:
        conn.execute("""INSERT INTO ml_notificaciones (topic, resource, recibido_en, cuenta_ml)
            VALUES ('post_purchase','/post-purchase/v1/claims/104','2026-10-16','principal')""")
    sync_ml._finalizar(0, __import__('datetime').datetime.utcnow(),
                      {"ordenes_vistas": 0, "ventas_upsert": 0, "envios_upsert": 0, "errores": 0}, "principal")
    with get_db() as conn:
        assert conn.execute("SELECT procesada FROM ml_notificaciones WHERE resource LIKE '%104'").fetchone()[0] == 0
    book = Workbook()
    book.active.append(["# de venta", "# de albarán", "Fecha de albarán"])
    book.active.append(["order-1", "ALB-1", __import__('datetime').date(2026, 10, 15)])
    excel = Path(tmp.name) / "con_fecha.xlsx"
    book.save(excel)
    assert parse_albaran(excel)["por_num_venta"] == 1
    book = Workbook()
    book.active.append(["# de venta", "# de albarán"])
    book.active.append(["order-1", "ALB-2"])
    excel = Path(tmp.name) / "sin_fecha.xlsx"
    book.save(excel)
    assert parse_albaran(excel)["por_num_venta"] == 1
    with get_db() as conn:
        row = conn.execute("SELECT albaran, fecha_albaran FROM ventas_ml WHERE num_venta='order-1'").fetchone()
        assert row["albaran"] == "ALB-2" and row["fecha_albaran"] == "2026-10-15"
        other = conn.execute("SELECT albaran, fecha_albaran FROM ventas_ml WHERE cuenta_ml='secundaria'").fetchone()
        assert other["albaran"] is None and other["fecha_albaran"] is None
    book = Workbook()
    book.active.append(["# de venta", "# de albarán", "Fecha de albarán"])
    book.active.append(["pack-1", "ALB-3", __import__('datetime').date(2026, 10, 16)])
    excel = Path(tmp.name) / "carrito.xlsx"
    book.save(excel)
    assert parse_albaran(excel)["por_pack_id"] == 1
    with get_db() as conn:
        own = conn.execute("SELECT albaran, fecha_albaran FROM ventas_ml WHERE num_venta='order-1'").fetchone()
        other = conn.execute("SELECT albaran, fecha_albaran FROM ventas_ml WHERE cuenta_ml='secundaria'").fetchone()
        assert own["albaran"] == "ALB-3" and own["fecha_albaran"] == "2026-10-16"
        assert other["albaran"] is None and other["fecha_albaran"] is None
    assert sync_devoluciones.reconciliar("principal") == 0
    rango = next(p["range"] for p in search_params if "range" in p)
    assert ",before:" in rango and ".000+00:00" in rango
    rango_params = next(p for p in search_params if "range" in p)
    assert rango_params["type"] == "returns"
    print("✅ postventa: duplicados, reintento 403/429, reclamo sin devolución, cruce exacto, SKU múltiples, plazo, filtros, marcas, evidencia y permisos")
finally:
    ml_client.get, ml_client.get_opcional = old_get, old_optional
    app.dependency_overrides.clear()
    tmp.cleanup()
