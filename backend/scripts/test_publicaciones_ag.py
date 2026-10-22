"""Regresión del flujo ARGENPARTS de publicaciones masivas."""
import os
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from services.generador_plantilla import COLUMNAS, ConfiguracionProveedor, escribir_xlsx, generar_filas_con_reporte
from services.imagenes_ag import asignar_imagenes_ag, leer_imagenes_ag
from services.parser_catalogo import cruzar_variantes, leer_catalogo_detallado, leer_publicaciones, normalizar_titulo
from services.perfiles_catalogo import perfil_de
from services.validacion_imagenes import _es_url_permitida

ok = bad = 0
def check(nombre, condicion):
    global ok, bad
    print(("  ✓ " if condicion else "  ✗ ") + nombre)
    ok += bool(condicion); bad += not bool(condicion)

headers = ["Producto", "Código", "SYD", "GROB", "KYB", "BOGE", "VAZLO", "Marca", "Aplicación", "Desde", "Hasta", "Posición", "Distribuidor", "stock"]
wb = openpyxl.Workbook(); ws = wb.active; ws.append(headers)
ws.append(["Amortiguador delantero", 1234, "S1", "", "K1", "", "V1", "Ford", "Focus (Europa)", 2012, 2015, "Delantero izquierdo", 100, 3])
ws.append(["Amortiguador delantero", "1234", "S2", "", "K1", "", "V1", "Ford", "Focus", 2016, 2018, "Delantero derecho", 100, 3])
ws.append(["Amortiguador", "BAD", "", "", "", "", "", "Ford", "Fiesta", 2020, 2019, "Delantero", 50, 1])
f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); f.close(); wb.save(f.name)
r = leer_catalogo_detallado(f.name, perfil_de("AG")); os.unlink(f.name)
check("detecta master AG por encabezados", r.formato == "master_ag")
check("Código numérico conserva SKU textual", r.piezas[0]["clave"] == "1234")
check("consolida equivalencias y compatibilidades", len(r.piezas[0]["compatibilidades"]) == 2 and r.piezas[0]["equivalencias"]["SYD"] == ["S1", "S2"])
check("años invertidos se excluyen y reportan", r.compatibilidades_invalidas == 1 and r.errores)
filas, reporte = generar_filas_con_reporte(r.piezas, ConfiguracionProveedor(codigo_bodega="AG", marca="ARGENPARTS"))
check("precio usa costo antes de IVA y genera variantes", len(filas) == 6 and filas[0].precio is not None)
focus_izquierdo = [x for x in filas if x.sku == "1234" and x.compatibilidad["posicion"] == "Delantero izquierdo"]
check("rango AG genera rango completo y dos bloques balanceados", len({x.titulo for x in focus_izquierdo}) == 3 and
      any("2012-2015" in x.titulo for x in focus_izquierdo) and
      any("2012 2013" in x.titulo for x in focus_izquierdo) and
      any("2014 2015" in x.titulo for x in focus_izquierdo))
check("cada título agrupado aparece una sola vez", len(focus_izquierdo) == len({x.titulo for x in focus_izquierdo}) == 3)
check("títulos AG no exceden 60 y conservan para", all(len(x.titulo) <= 60 and " para " in x.titulo for x in filas))
check("descripción conserva notas parentéticas", "(Europa)" in filas[0].descripcion)

# La identidad anual es SKU + producto + posición + marca + aplicación + año,
# aunque dos filas equivalentes lleguen al generador sin pasar por el parser.
pieza_duplicada = {
    "formato": "master_ag", "clave": "DUP-AG", "producto": "Amortiguador",
    "linea": "Amortiguador", "stock": 1, "costo": 100,
    "compatibilidades": [
        {"marca": "Ford", "aplicacion": "Focus", "posicion": "Delantero", "inicio": 2018, "fin": 2019, "fila": 1},
        {"marca": "Ford", "aplicacion": "Focus", "posicion": "Delantero", "inicio": 2018, "fin": 2019, "fila": 2},
    ], "equivalencias": {},
}
filas_dup, reporte_dup = generar_filas_con_reporte([pieza_duplicada], ConfiguracionProveedor("AG", marca="ARGENPARTS"))
check("deduplica por SKU + compatibilidad + título agrupado", len(filas_dup) == 3 and reporte_dup["deduplicadas"] == 3)

# El reporte ML puede conservar títulos AG históricos con rango. Debe cubrir
# cada año sólo cuando todos los campos de la compatibilidad siguen presentes.
pieza_cruce = {
    "formato": "master_ag", "clave": "AG-CRUCE", "producto": "Amortiguador delantero",
    "linea": "Amortiguador delantero", "stock": 1, "costo": 100, "equivalencias": {},
    "compatibilidades": [
        {"marca": "Ford", "aplicacion": "Focus (Europa)", "posicion": "Delantero izquierdo", "inicio": 2012, "fin": 2015, "fila": 1},
        {"marca": "Ford", "aplicacion": "Focus (Europa)", "posicion": "Delantero derecho", "inicio": 2013, "fin": 2013, "fila": 2},
        {"marca": "Ford", "aplicacion": "Fiesta", "posicion": "Delantero izquierdo", "inicio": 2013, "fin": 2013, "fila": 3},
    ],
}
filas_cruce, _ = generar_filas_con_reporte([pieza_cruce], ConfiguracionProveedor("AG", marca="ARGENPARTS"))
objetivo = [x for x in filas_cruce if x.compatibilidad["aplicacion"] == "Focus (Europa)" and
            x.compatibilidad["posicion"] == "Delantero izquierdo"]
pub = openpyxl.Workbook(); ps = pub.active; ps.append(["col"] * 25)
def pubrow(sku, titulo):
    row = [None] * 25; row[1] = titulo; row[16] = sku; ps.append(row)
pubrow("AG-CRUCE", "Amort. Del. Izq. P/ Ford Focus 2012-2015")
rp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); rp.close(); pub.save(rp.name)
publicados = leer_publicaciones(rp.name); os.unlink(rp.name)
cruzado = cruzar_variantes(filas_cruce, publicados)
check("un rango histórico AG cubre rango y bloques contenidos", len([x for x in cruzado["existentes"] if x.compatibilidad["posicion"] == "Delantero izquierdo"]) == 3)
check("otro vehículo o posición con el mismo SKU queda pendiente", len(cruzado["pendientes"]) == 2 and
      {x.compatibilidad["aplicacion"] for x in cruzado["pendientes"]} == {"Focus (Europa)", "Fiesta"})

pub_ambiguo = openpyxl.Workbook(); pam = pub_ambiguo.active; pam.append(["col"] * 25)
row = [None] * 25; row[1] = "Amort. Del. Izq. Der. P/ Ford Focus 2013"; row[16] = "AG-CRUCE"; pam.append(row)
ram = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); ram.close(); pub_ambiguo.save(ram.name)
publicados_ambiguos = leer_publicaciones(ram.name); os.unlink(ram.name)
anio_posiciones = [x for x in filas_cruce if x.aplicacion == "2013" and
                   x.compatibilidad["aplicacion"] == "Focus (Europa)"]
check("una coincidencia AG ambigua queda pendiente", not cruzar_variantes(anio_posiciones, publicados_ambiguos)["existentes"])

pub_anual = openpyxl.Workbook(); pa = pub_anual.active; pa.append(["col"] * 25)
row = [None] * 25; row[1] = "Amortiguador Delantero Izquierdo P/ Ford Focus 2013"; row[16] = "AG-CRUCE"; pa.append(row)
ra = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); ra.close(); pub_anual.save(ra.name)
publicados_anual = leer_publicaciones(ra.name); os.unlink(ra.name)
cruzado_anual = cruzar_variantes(objetivo, publicados_anual)
check("una publicación anual AG no cubre un rango o bloque incompleto", not cruzado_anual["existentes"] and len(cruzado_anual["pendientes"]) == 3)

fotos = openpyxl.Workbook(); hoja = fotos.active; hoja.append(["Servidor", "Carpeta", "Nombre", "Extensión"])
for n in ("1234", "1234_2", "1234-3", "1234_11", "KIT 1234", "SINMATCH"):
    hoja.append(["https://apimagenes.s3.amazonaws.com", "AG", n, "jpg"])
ff = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); ff.close(); fotos.save(ff.name)
galeria, resumen = leer_imagenes_ag(ff.name, ["1234"]); os.unlink(ff.name)
check("fotos: principal, guion bajo y guion respetan orden", galeria["1234"][:3] == ["https://apimagenes.s3.amazonaws.com/AG/1234.jpg", "https://apimagenes.s3.amazonaws.com/AG/1234_2.jpg", "https://apimagenes.s3.amazonaws.com/AG/1234-3.jpg"])
check("fotos omite kit, fuera de límite y SKU inexistente", resumen["urls_omitidas"] == 3)

# Encabezados reales de FOTOS AG.xlsx. La primera hoja trae ``2`` en vez de
# ``URL servidor``: URL FINAL debe funcionar aunque ese campo no sea usable.
fotos_reales = openpyxl.Workbook(); hoja = fotos_reales.active
hoja.append([2, "Nombre de la carpeta", "Nombre de la foto", "Formato de la foto", "URL FINAL"])
hoja.append(["https://servidor-antiguo.example", "confort2024", "1234", ".png",
             "https://apimagenes.s3.amazonaws.com/confort2024/1234.jpg"])
hoja.append(["https://servidor-antiguo.example", "confort2024", "1234_2", ".png",
             "https://apimagenes.s3.amazonaws.com/confort2024/1234_2.jpg"])
ff = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); ff.close(); fotos_reales.save(ff.name)
galeria_real, resumen_real = leer_imagenes_ag(ff.name, ["1234"]); os.unlink(ff.name)
check("fotos reconoce encabezados reales de AG y URL FINAL", galeria_real["1234"][:2] == [
    "https://apimagenes.s3.amazonaws.com/confort2024/1234.jpg",
    "https://apimagenes.s3.amazonaws.com/confort2024/1234_2.jpg",
] and resumen_real["urls_cruzadas"] == 2)
check("fotos AG sólo permiten el host HTTPS exacto", _es_url_permitida(
    "https://apimagenes.s3.amazonaws.com/AG/1234.jpg", ("apimagenes.s3.amazonaws.com",)) and not _es_url_permitida(
    "https://cdn.apimagenes.s3.amazonaws.com/AG/1234.jpg", ("apimagenes.s3.amazonaws.com",)))
asignar_imagenes_ag(filas, galeria)
out = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); out.close(); escribir_xlsx(filas, ConfiguracionProveedor(codigo_bodega="AG", marca="ARGENPARTS"), out.name)
salida = openpyxl.load_workbook(out.name, data_only=True); hoja = salida.active
check("XLSX conserva 37 columnas, Cantidad y bodega AG", [c.value for c in hoja[1]] == COLUMNAS and hoja.cell(2, 5).value == 3 and hoja.cell(2, 31).value == 3 and hoja.cell(2, 15).value)
salida.close(); os.unlink(out.name)

# Master de combos AG: aplicación/años sin posición y componentes consolidados.
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["Codigo", "Producto", "Aplicación", "Año desde", "Año hasta", "Componente 1", "Componente 2", "Stock", "Precio"])
ws.append(["AG-COMBO", "Kit distribución", "Ford Focus", 2018, 2020, "A-1", "B-2", 5, 200])
ws.append(["AG-COMBO", "Kit distribución", "Ford Focus", 2018, 2020, "B-2", "C-3", 5, 200])
ws.append(["AG-COMBO", "Kit distribución", "Ford Fiesta", 2021, 2021, "D-4", None, 5, 200])
ws.append(["BAD-STOCK", "Kit", "Ford Focus", 2020, 2020, "x", None, "mal", 100])
ws.append(["BAD-COST", "Kit", "Ford Focus", 2020, 2020, "x", None, 2, 100])
ws.append(["BAD-COST", "Kit", "Ford Fiesta", 2020, 2020, "y", None, 2, 101])
ws.append(["BAD-YEAR", "Kit", "Ford Focus", 2022, 2020, "x", None, 1, 100])
f = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); f.close(); wb.save(f.name)
rc = leer_catalogo_detallado(f.name, perfil_de("AG")); os.unlink(f.name)
check("detecta catálogo de combos AG y consolida componentes", rc.formato == "master_ag" and rc.piezas[0]["tipo_ag"] == "combos" and rc.piezas[0]["componentes"] == ["A-1", "B-2", "C-3", "D-4"])
check("combo consolida aplicaciones repetidas y reporta años inválidos", len(rc.piezas[0]["compatibilidades"]) == 2 and rc.compatibilidades_invalidas == 1)
check("excluye SKU con stock inválido o costos distintos", {p["clave"] for p in rc.piezas} == {"AG-COMBO"} and any("Stock" in e["motivo"] for e in rc.errores) and any("Precio diferente" in e["motivo"] for e in rc.errores))
wb = openpyxl.Workbook(); wh = wb.active; wh.append(["Combos por aplicaciones AG"]); wh.append([None]); wh.append(["Codigo", "Producto", "Aplicación", "Año desde", "Año hasta", "Componente 1", "Stock", "Precio"]); wh.append(["SKU-ROW3", "Kit", "Ford Focus", 2019, 2020, "P-1", 2, 125])
fh = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); fh.close(); wb.save(fh.name)
rh = leer_catalogo_detallado(fh.name, perfil_de("AG")); os.unlink(fh.name)
check("detecta encabezados del master AG debajo de filas de título", len(rh.piezas) == 1 and rh.piezas[0]["clave"] == "SKU-ROW3")
wb = openpyxl.Workbook(); wm = wb.active; wm.append(["Codigo", "Producto", "Aplicación", "Año desde", "Año hasta", "Stock", "Precio"]); wm.append(["MISSING", "Kit", "Ford Focus", 2019, 2020, None, None])
fm = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); fm.close(); wb.save(fm.name)
rm = leer_catalogo_detallado(fm.name, perfil_de("AG")); os.unlink(fm.name)
check("reporta Stock y Precio vacíos por separado", not rm.piezas and {e["motivo"] for e in rm.errores} == {"Stock vacío dentro del SKU", "Precio vacío dentro del SKU"})
fc, _ = generar_filas_con_reporte(rc.piezas, ConfiguracionProveedor("AG", marca="ARGENPARTS"))
check("títulos de combos omiten posición y la descripción lista componentes", all("para Ford" in x.titulo and len(x.titulo) <= 60 and "delantero" not in x.titulo.lower() for x in fc) and "A-1 | B-2 | C-3 | D-4" in fc[0].descripcion)
fotos_combo = openpyxl.Workbook(); sc = fotos_combo.active; sc.append(["Servidor", "Carpeta", "Nombre", "Extensión"])
sc.append(["https://apimagenes.s3.amazonaws.com", "AG", "AG-COMBO", "jpg"])
sc.append(["https://apimagenes.s3.amazonaws.com", "AG", "AG-COMBO_2", "jpg"])
ff = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); ff.close(); fotos_combo.save(ff.name)
gc, _ = leer_imagenes_ag(ff.name, ["AG-COMBO"], exact_only=True); os.unlink(ff.name)
check("fotos de combos cruzan sólo por código completo exacto", len(gc["AG-COMBO"]) == 10 and bool(gc["AG-COMBO"][0]) and not gc["AG-COMBO"][1])
outc = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); outc.close()
asignar_imagenes_ag(fc, gc)
escribir_xlsx(fc, ConfiguracionProveedor("AG", marca="ARGENPARTS"), outc.name)
oc = openpyxl.load_workbook(outc.name, data_only=True); hc = oc.active
check("combo escribe precio, stock, bodega AG e imágenes en las 37 columnas", [c.value for c in hc[1]] == COLUMNAS and hc.cell(2, 5).value == 5 and hc.cell(2, 31).value == 5 and hc.cell(2, 3).value is not None and hc.cell(2, 15).value)
oc.close(); os.unlink(outc.name)
cruce_combo = cruzar_variantes(fc, {("AG-COMBO", normalizar_titulo("Kit distribución para Ford Focus 2018-2020"))})
check("publicación histórica AG por SKU y compatibilidad cubre variante combo", len(cruce_combo["existentes"]) >= 1)
print(f"\n{ok} pasaron · {bad} fallaron")
raise SystemExit(bool(bad))
